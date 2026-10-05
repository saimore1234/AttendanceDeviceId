"""Excel (.xlsx) import adapter - reads the file attached on the
Attendance Device (import_file) using the column mapping configured
there. First row is treated as the header row."""

from __future__ import annotations

import frappe
from openpyxl import load_workbook

from attendance_device_integration.device_adapters.base import AdapterError
from attendance_device_integration.device_adapters.file_import_base import FileImportAdapterBase


class ExcelAdapter(FileImportAdapterBase):
	adapter_name = "Excel Import"

	def _rows(self):
		file_url = self.device.get("import_file")
		try:
			file_path = frappe.utils.file_manager.get_file_path(file_url)
		except Exception as e:
			raise AdapterError(f"Could not resolve import file: {e}", error_code="INVALID_CONFIG")

		try:
			wb = load_workbook(file_path, read_only=True, data_only=True)
		except Exception as e:
			raise AdapterError(f"Could not read Excel file: {e}", error_code="INVALID_CONFIG")

		ws = wb.active
		rows = ws.iter_rows(values_only=True)
		try:
			headers = next(rows)
		except StopIteration:
			return

		for row in rows:
			if all(v is None for v in row):
				continue
			yield dict(zip(headers, row))
